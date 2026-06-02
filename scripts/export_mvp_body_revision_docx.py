from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUT = Path("docs/mvp_body_revision_and_appendix.docx")
SRC = Path("docs/mvp_body_revision_and_appendix_source.md")


def set_run_font(run, size: float = 10.5, bold: bool | None = None) -> None:
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_para(doc: Document, parts: list[tuple[str, bool]], style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    for text, bold in parts:
        run = paragraph.add_run(text)
        set_run_font(run, bold=bold)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading("", level=level)
    run = paragraph.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        set_run_font(run, size=9.5, bold=True)
        set_cell_shading(cell, "D9EAF7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = ""
            paragraph = cells[idx].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            for line_idx, line in enumerate(str(value).split("\n")):
                if line_idx:
                    paragraph.add_run("\n")
                run = paragraph.add_run(line)
                set_run_font(run, size=9.3)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    for table_row in table.rows:
        for idx, width in enumerate(widths):
            table_row.cells[idx].width = Inches(width)
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    style = doc.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    style.font.size = Pt(10.5)

    add_heading(doc, "MVP 반영 본문 수정안 및 부록 초안", 1)
    add_para(doc, [("※ 굵게 표시한 부분은 기존 수정본 대비 보강하거나 표현 강도를 조정한 부분입니다.", True)])

    add_heading(doc, "마. 군 단독망 적용을 위한 로컬 추론 가능성의 예비 검토", 2)
    add_para(
        doc,
        [
            (
                "군 조직의 실제 적용환경을 고려하면, 공개 API를 호출하는 방식만으로는 실질적 활용성을 충분히 설명하기 어렵다. "
                "군은 단독망 환경에서의 운용이 일반적이며, 중앙 서버 기반 최신 LLM 또는 고성능 GPU 서버를 안정적으로 구축·운영하는 데 현실적 제약이 존재하기 때문이다. ",
                False,
            ),
            (
                "본 예비 시연의 목적은 모델 간 우열 비교가 아니라, 제안한 버전관리형 지식기반과 검색 구조가 단독망 또는 로컬 환경에서도 작동 가능한지를 확인하는 데 있다.",
                True,
            ),
        ],
    )
    add_para(
        doc,
        [
            (
                "이를 위해 본 연구에서는 공개 API 기반 MVP와 별도로 오픈소스 모델을 활용한 로컬 구동을 예비적으로 시연하였다. "
                "시연은 단순한 단독 추론이 아니라, ",
                False,
            ),
            ("정규화된 법령 자료, 벡터DB 기반 검색, 근거 카드, LLM 요약, 검증 및 fallback 구조를 결합한 상태", True),
            (
                "에서 수행하였다. 본문에서 다루는 핵심 코퍼스는 「군인의 지위 및 복무에 관한 기본법」, 시행령, 시행규칙이며, ",
                False,
            ),
            ("군인복무규율은 2016년 체계 전환을 설명하기 위한 연혁 보조자료로 한정하여 포함하였다.", True),
        ],
    )
    add_para(
        doc,
        [
            (
                "MVP 기능과 질의 유형의 연결은 다음과 같이 설정하였다. 단순한 현행 조문 확인은 검색형 기능을, "
                "개정 이유 및 조문 간 연결 질문은 설명·비교형 기능을, 실제 업무 처리 전에 추가로 확인해야 할 사항을 묻는 질문은 "
                "실무 안내형 기능을 점검하기 위한 용도로 사용하였다. ",
                False,
            ),
            (
                "다만 본 절의 결과는 제한된 질문과 특정 장비 환경에서 수행한 예비 시연 결과이며, 본격적인 성능평가나 모델 벤치마크로 일반화하지 않는다.",
                True,
            ),
        ],
    )
    add_table(
        doc,
        ["MVP 기능", "대표 질문 유형", "확인하려는 사항"],
        [
            ["검색형", "현행 기본법 제45조 신고자 보호 규정 설명", "조문번호와 현행 문언을 올바르게 찾아 근거 카드와 연결하는지 확인"],
            [
                "설명·비교형",
                "기본법 제43조, 제44조, 제45조의 신고의무·비밀보장·신고자 보호 연결 구조 설명",
                "관련 조문을 시간·체계 맥락 속에서 구분해 설명하는지 확인",
            ],
            ["실무 안내형", "휴가 종류 또는 신고 사안 검토 시 추가 확인사항 제시", "현행 조문 외에 시행령, 부대 지침, 원문 확인 필요성을 구분해 제시하는지 확인"],
        ],
        [1.15, 2.75, 3.65],
    )
    add_para(
        doc,
        [
            ("대표 예비 시연 결과, 두 모델 모두 벡터DB 기반 검색과 결합했을 때 기초적인 근거 기반 응답을 생성할 수 있었다. ", False),
            ("그러나 단순질문에서는 비교적 안정적이었던 반면, 복합 비교형 질문에서는 해석 오류와 비교 오류가 관찰되어 추가 검증의 필요성이 확인되었다.", True),
        ],
    )
    add_table(
        doc,
        ["모델", "질문 유형", "응답 시간/토큰", "관찰된 강점", "관찰된 한계"],
        [
            ["EXAONE-4.0-1.2B", "단순 조문 설명", "1.8초 / 616토큰", "응답 속도가 빠르고 현행 조문 요약 가능", "일부 조문 의미를 뭉뚱그려 설명할 가능성"],
            ["GPT-OSS-20B", "단순 조문 설명", "33.9초 / 1,277토큰", "비교적 상세한 서술 가능", "응답 시간이 길고 reasoning 토큰 소모가 큼"],
            ["EXAONE-4.0-1.2B", "복합 비교형 질문", "15.0초 / 780토큰", "과거 규정과 현행 조문의 연결 구조를 대체로 제시", "핵심 결론에서 조문 기능을 일부 포괄적으로 묶는 경향"],
            ["GPT-OSS-20B", "복합 비교형 질문", "29.5초 / 1,373토큰", "설명량은 충분하나 검증 필요", "연혁 자료와 모순되는 응답이 관찰됨"],
        ],
        [1.35, 1.35, 1.35, 2.0, 2.1],
    )
    add_para(
        doc,
        [
            (
                "이 결과는 본 연구가 제안하는 구조가 반드시 중앙 GPU 서버만을 전제로 하지 않으며, 적절한 자료 구조와 검색정책이 마련될 경우 "
                "사무실·부서 단위 로컬 환경에서도 검토 가능한 대안이 될 수 있음을 시사한다. ",
                False,
            ),
            ("다만 이는 운용 가능성의 기초 자료일 뿐이며, 장시간 연속 사용, 동시 접속, 대규모 내부자료 탑재, 다양한 복합질의에 대한 안정성은 후속 검증이 필요하다.", True),
        ],
    )
    add_para(
        doc,
        [
            (
                "또한 비교 과정에서 확인된 중요한 점은 모델 자체의 크기보다 지식기반의 구조와 검색 품질이 응답의 유용성을 크게 좌우했다는 것이다. "
                "외부 문서 연결이 없는 기본 모델은 질문의 맥락이나 개정 경위를 안정적으로 설명하는 데 한계를 보였으나, 구조화된 자료와 검색정책을 결합한 경우에는 "
                "상대적으로 일관된 근거 제시가 가능하였다. ",
                False,
            ),
            ("따라서 군 단독망 적용을 논의할 때에는 더 큰 모델을 도입하는 것만큼이나 문서관리체계, 메타데이터 설계, 검색정책, 근거 중심 요약 구조를 정교화하는 것이 중요하다.", True),
        ],
    )

    add_heading(doc, "바. MVP의 한계와 시사점", 2)
    add_para(doc, [("본 MVP는 다음과 같은 한계를 가진다.", False)])
    limitations = [
        ("첫째, 공개 법령 자료 중심으로 구성되었기 때문에 실제 군 조직에서 중요한 내부 예규, 지침, 공문, 해석 메모 등을 포함하지 못하였다.", False),
        ("둘째, 시점 인지형 검색을 지향했지만 모든 시점의 규정 상태와 실제 적용 맥락을 완전하게 복원한 것은 아니다.", False),
        ("셋째, 생성형 AI의 응답은 근거를 요약·정리하는 보조기능이며 최종 판단을 대체할 수 없다.", False),
        ("넷째, 문서 반입 절차, 접근권한 통제, 모델·문서 업데이트 검증, 감사로그 등 운영 거버넌스 요소는 본 MVP 범위에 포함되지 않았다.", True),
        ("다섯째, 로컬 LLM 예비 시연은 제한된 질의와 특정 장비 환경에서 수행되었으므로 모든 일반 PC 환경에서의 안정적 운용으로 일반화할 수 없다.", True),
    ]
    for text, bold in limitations:
        add_para(doc, [("· ", False), (text, bold)])
    add_para(
        doc,
        [
            (
                "특히 복합 비교형 질문에서는 한계가 더 분명히 드러났다. EXAONE-4.0-1.2B는 과거 규정과 현행 조문을 대체로 연결했으나 "
                "핵심 결론에서 일부 조문 기능을 포괄적으로 묶는 경향을 보였고, GPT-OSS-20B는 일부 응답에서 동일 화면의 연혁자료와 모순되는 설명을 제시하였다. ",
                False,
            ),
            ("이는 내부 자동검사 지표가 긍정적으로 표시되더라도, 복합 비교형 질문에서는 별도 전문가 검토와 근거 대조가 필요함을 보여준다.", True),
        ],
    )
    add_para(doc, [("그럼에도 본 시범 구현은 다음과 같은 시사점을 가진다.", False)])
    implications = [
        ("첫째, 조직기억 아카이브와 시점 인지형 검색이라는 설계 방향이 추상적 개념에 머무르지 않고 실제 시스템 구조와 사용자 인터페이스로 구체화될 수 있음을 보여준다.", False),
        ("둘째, 규정업무 지원에서 응답의 유용성은 모델 규모 자체보다 구조화된 지식기반, 메타데이터, 검색정책, 안전한 요약생성 구조에 크게 좌우될 수 있다.", True),
        ("셋째, 군의 단독망 특성을 고려할 때 공개 API 기반 시연뿐 아니라 사무실·부서 단위의 로컬 구동 구조도 후속 단계에서 검토할 수 있는 현실적 대안임을 보여준다.", False),
    ]
    for text, bold in implications:
        add_para(doc, [("· ", False), (text, bold)])
    add_para(
        doc,
        [
            ("따라서 본 MVP의 의미는 특정 모델의 성능을 입증하는 데 있지 않다. ", False),
            ("핵심은 개정 이력과 근거 문서를 구조화하여 보존하고, 그 구조를 검색·요약 계층이 활용할 때 조직기억의 단절을 줄일 수 있는지 확인했다는 점에 있다.", True),
        ],
    )

    add_heading(doc, "부록. MVP 구현 구조 및 예비 시연 결과", 1)
    add_heading(doc, "부록 1. MVP 구현 구조 및 로직", 2)
    add_para(
        doc,
        [
            (
                "본 부록은 MVP를 재현 가능한 수준에서 설명하기 위한 것이며, 특정 도구 사용법을 안내하는 개발 매뉴얼이 아니라 연구 시연의 구조와 판단 기준을 제시하는 데 목적이 있다.",
                True,
            )
        ],
    )
    add_table(
        doc,
        ["단계", "처리 내용", "산출물 또는 확인 사항"],
        [
            ["1. 원천자료 수집", "공개 법령, 시행령, 시행규칙, 개정이유, 연혁자료, 필요한 경우 군인복무규율 연혁자료를 수집", "raw 자료와 출처 URL 보존"],
            ["2. 정규화", "PDF/HTML/TXT 자료를 조문 단위 또는 개정이유 단위로 정리하고 메타데이터 부여", "processed JSONL"],
            ["3. 색인화", "정규화 문서를 임베딩하여 벡터DB에 저장하고, 법령명·조문번호·자료유형 기반 필터 검색을 병행", "Chroma 또는 JSON fallback store"],
            ["4. 질의 분류", "현행 조문 확인, 개정이유 설명, 비교, 실무 참고, 복합질문으로 분류", "intent 및 route rationale"],
            ["5. 검색 및 답변 구성", "근거 문서를 우선 검색하고 deterministic answer builder로 기본 답변 골격 작성", "근거 카드와 기본 답변"],
            ["6. LLM 요약", "LLM은 전체 답변 생성자가 아니라 핵심 결론과 짧은 해석을 보조적으로 생성", "요약 생성 또는 fallback"],
            ["7. 검증 및 표시", "짧은 답변, 근거 없는 추가, 연도·조문 불일치 등을 점검하고 필요 시 근거 중심 답변으로 전환", "사용자 답변, debug 정보, DOCX export"],
        ],
        [1.2, 3.3, 3.0],
    )
    add_para(
        doc,
        [
            (
                "답변 화면은 본문 답변과 근거 카드를 분리한다. 본문 답변은 핵심 결론, 세부 정리, 실무 참고, 근거 안내로 구성하고, "
                "근거 카드는 법령명, 자료유형, 조문번호 또는 범위, 시행일, 원문 링크를 제시한다.",
                False,
            )
        ],
    )

    add_heading(doc, "부록 2. 예비 시연 환경 및 질문 세트", 2)
    add_table(
        doc,
        ["항목", "내용", "비고"],
        [
            ["CPU", "기재 필요", "시연 장비 사양 확인 후 입력"],
            ["GPU / VRAM", "기재 필요", "로컬 추론 가능성 판단의 핵심 정보"],
            ["RAM", "기재 필요", "모델 로딩과 벡터DB 운용 영향"],
            ["OS", "Windows 환경", "정확한 버전 기재 권장"],
            ["추론 도구", "LM Studio", "도구 자체보다 로컬 추론 구조가 핵심"],
            ["사용 모델", "EXAONE-4.0-1.2B, GPT-OSS-20B", "모델 간 우열 비교가 아니라 운용 가능성 확인 목적"],
            ["임베딩 모델", "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2", "한국어 질의와 법령 문서 검색에 활용"],
            ["벡터DB", "Chroma 기반 색인 및 JSON fallback", "검색 실패 시 보수적 근거 답변 유지"],
            ["생성 설정", "temperature 0.0, 모델별 max token 제한", "환각 및 장문 실패를 줄이기 위한 보수적 설정"],
        ],
        [1.6, 3.0, 3.0],
    )
    add_table(
        doc,
        ["구분", "질문 예시", "검증 목적"],
        [
            ["현행 조문 설명 1", "현행 기본법 제45조 신고자 보호 규정을 설명해줘.", "조문 식별 정확성"],
            ["현행 조문 설명 2", "현행 시행령 기준으로 군인의 휴가 종류를 조문번호와 함께 나열해줘.", "현행 조문 검색과 요약"],
            ["특정 시점 확인 1", "2016년 기본법 제정 당시 신고의무와 신고자 보호 취지를 설명해줘.", "시점 일치성"],
            ["특정 시점 확인 2", "군인복무규율 폐지 전후 사적 제재 금지 규정이 어떻게 연결되는지 설명해줘.", "연혁자료 활용"],
            ["개정 이유 설명 1", "기본법 제정 이유에서 군 내 기본권 침해 문제를 어떻게 설명했는지 정리해줘.", "개정 이유 근거 제시"],
            ["개정 이유 설명 2", "휴가·돌봄 관련 시행령 개정 이유와 현행 조문을 연결해 설명해줘.", "개정 취지와 현행 조문 연결"],
            ["신·구 비교 1", "기본법 제43조부터 제45조의 신고의무, 비밀보장, 신고자 보호를 연결해서 설명해줘.", "조문 관계 비교"],
            ["신·구 비교 2", "군인복무규율 제15조와 현행 기본법 제43조부터 제45조를 비교해줘.", "과거 규정과 현행 체계 비교"],
            ["실무 안내 1", "가혹행위 신고 사안을 검토할 때 어떤 조문과 원문을 추가 확인해야 하는지 알려줘.", "실무상 확인사항"],
            ["실무 안내 2", "휴가 제한 또는 보류를 판단할 때 기본법과 시행령에서 무엇을 함께 확인해야 하는지 알려줘.", "현행 조문과 하위 법령 연결"],
        ],
        [1.4, 4.0, 2.2],
    )

    add_heading(doc, "부록 3. 대표 응답 비교 및 관찰 결과", 2)
    add_table(
        doc,
        ["질문 유형", "모델", "결과 요약", "판정 초안"],
        [
            ["단순 조문 설명", "EXAONE-4.0-1.2B", "빠르게 조문 내용을 요약하였고 근거 카드와 연결되었다. 다만 일부 표현은 조문 기능을 넓게 묶는 경향이 있었다.", "부분정확"],
            ["단순 조문 설명", "GPT-OSS-20B", "상대적으로 상세한 설명을 생성하였으나 응답 시간이 길고 reasoning 토큰 사용량이 컸다.", "부분정확"],
            ["복합 비교형", "EXAONE-4.0-1.2B", "과거 군인복무규율과 현행 기본법 조문을 대체로 연결했으나 핵심 결론에서 일부 조문 의미가 포괄적으로 정리되었다.", "부분정확"],
            ["복합 비교형", "GPT-OSS-20B", "설명량은 충분했으나 일부 응답에서 연혁자료와 모순되는 내용이 관찰되었다.", "오류 포함"],
        ],
        [1.5, 1.5, 3.8, 1.0],
    )
    add_para(
        doc,
        [
            (
                "평가 기준은 조문 식별 정확성, 시점 일치 여부, 비교 정확성, 개정 이유 근거 유무, 실무상 확인사항의 적절성, 허위 생성 여부로 설정하는 것이 적절하다. "
                "가능하다면 규정 실무자와 법무 검토자가 각각 정확, 부분정확, 오류로 판정하는 간단한 전문가 검토 절차를 추가하는 것이 바람직하다.",
                True,
            )
        ],
    )

    add_heading(doc, "부록 4. 작성 및 운용상 제한사항", 2)
    appendix_limits = [
        "작은 모델은 단순 조문 확인에는 비교적 안정적이나, 복합 비교형 질문에서는 조문 관계를 뭉뚱그리거나 일부 근거와 모순되는 응답을 낼 수 있다.",
        "GPT-OSS-20B와 같은 reasoning 계열 모델은 내부 reasoning 토큰을 많이 사용하여 최종 답변이 짧아지거나 끊기는 현상이 발생할 수 있다.",
        "자동 validator가 통과하더라도 실질적 정확성이 보장되는 것은 아니므로, 고위험 질문은 담당자 검토 절차가 필요하다.",
        "벡터DB 검색 품질은 원천자료 정규화, 조문번호 메타데이터, 자료유형 구분, 검색어 확장 정책에 크게 의존한다.",
        "공개 법령 중심 MVP이므로 내부 예규, 훈령, 공문, 해석 메모, 반복 민원 대응 기준을 포함하지 못한다.",
        "군 내부 적용 시에는 문서 반입 승인, 보안등급 분류, 접근권한, 로그 보관, 개인정보 마스킹, 모델 업데이트 검증, 감사 체계를 별도로 설계해야 한다.",
    ]
    for item in appendix_limits:
        add_para(doc, [("· ", False), (item, False)])

    return doc


def write_source() -> None:
    SRC.write_text(
        """# MVP 반영 본문 수정안 및 부록 초안

굵게 표시한 부분은 기존 수정본 대비 보강하거나 표현 강도를 조정한 부분입니다.

## 본문 수정 방향

- 제목을 “군 단독망 적용을 위한 로컬 추론 가능성의 예비 검토”로 조정했습니다.
- 모델 성능 비교가 아니라 로컬·단독망 적용 가능성의 예비 검토임을 명확히 했습니다.
- 군인복무규율은 핵심 코퍼스가 아니라 2016년 체계 전환을 설명하기 위한 연혁 보조자료로 한정했습니다.
- 그림 2·3은 본문보다 부록 배치가 적절하므로, 본문에는 기능 매핑표와 결과 요약표 중심으로 정리했습니다.
- 부록에는 MVP 구현 구조, 시연 환경, 질문 세트, 대표 응답 비교, 작성 및 운용상 제한사항을 넣었습니다.
""",
        encoding="utf-8",
    )


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUT)
    write_source()
    print(OUT.resolve())
    print(SRC.resolve())


if __name__ == "__main__":
    main()
