from __future__ import annotations

from pathlib import Path
import re

from docx import Document


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_PATH = PROJECT_ROOT / "docs" / "chapters_3_to_5_draft_source.md"
OUTPUT_PATH = PROJECT_ROOT / "docs" / "chapters_3_to_5_draft.docx"


def add_bold_runs(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def build_docx(source_text: str) -> Document:
    document = Document()

    for raw_line in source_text.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if not stripped:
            document.add_paragraph("")
            continue
        if stripped.startswith("# "):
            document.add_heading(stripped[2:].strip(), level=0)
            continue
        if stripped.startswith("## "):
            document.add_heading(stripped[3:].strip(), level=1)
            continue
        if stripped.startswith("### "):
            document.add_heading(stripped[4:].strip(), level=2)
            continue
        if stripped.startswith("#### "):
            document.add_heading(stripped[5:].strip(), level=3)
            continue
        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_bold_runs(paragraph, stripped[2:].strip())
            continue

        paragraph = document.add_paragraph()
        add_bold_runs(paragraph, stripped)

    return document


def main() -> None:
    source_text = SOURCE_PATH.read_text(encoding="utf-8")
    document = build_docx(source_text)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
